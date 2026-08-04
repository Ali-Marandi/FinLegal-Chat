import os
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
import pandas as pd

class DocumentAutomator:
    @staticmethod
    def create_legal_report(title: str, content: dict, output_path: str):
        """Generates a professional Word document report."""
        doc = Document()
        doc.add_heading(title, 0)
        
        for section, text in content.items():
            doc.add_heading(section, level=1)
            p = doc.add_paragraph(text)
            p.style.font.size = Pt(11)
            
        doc.save(output_path)
        return output_path

    @staticmethod
    def create_financial_spreadsheet(data: dict, output_path: str):
        """Generates a professional Excel spreadsheet from financial data."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Financial Analysis"
        
        # Assume data is a dict with 'labels' and 'values'
        ws.append(["Category", "Value"])
        for label, value in zip(data.get('labels', []), data.get('values', [])):
            ws.append([label, value])
            
        wb.save(output_path)
        return output_path
